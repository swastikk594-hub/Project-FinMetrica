from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION

OUT = "docs/student_methodology_paper/quantfinance_edu_student_methodology.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_text(p, text, bold=False, italic=False, size=None):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.name = "Aptos"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    if size:
        r.font.size = Pt(size)
    return r


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    add_text(p, text, bold=True, size=15 if level == 1 else 12)
    return p


def para(doc, text, lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.15
    if lead:
        add_text(p, lead, bold=True)
    add_text(p, text)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    add_text(p, text)
    return p


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]
        c.width = Inches(widths[i])
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(c, "1F4E79")
        borders(c)
        set_cell_margins(c)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, h, bold=True, size=9)
        for run in p.runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            c = cells[i]
            c.width = Inches(widths[i])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx % 2 == 1:
                shade(c, "EEF5FA")
            borders(c)
            set_cell_margins(c)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_text(p, value, size=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.72)
sec.bottom_margin = Inches(0.72)
sec.left_margin = Inches(0.78)
sec.right_margin = Inches(0.78)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(10.5)
for s in ("Heading 1", "Heading 2"):
    styles[s].font.name = "Aptos Display"
    styles[s]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    styles[s]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    styles[s].font.color.rgb = RGBColor(0, 0, 0)

# Title page
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(100)
p.paragraph_format.space_after = Pt(16)
add_text(p, "How a Quantitative Investing Algorithm Makes Decisions", bold=True, size=25)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(22)
add_text(p, "A student friendly methodology paper for ages 13 to 18", italic=True, size=14)
for line in [
    "QuantFinance EDU Personal Project",
    "Purpose: explain the method, the checks, and the limits of a data based investing model",
    "This is an educational research project. It is not financial advice or a promise of profit.",
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    add_text(p, line, size=11)
doc.add_page_break()

heading(doc, "Summary")
para(doc, "This paper explains the method behind QuantFinance EDU, an algorithm that studies shares and exchange traded funds using historical information. Its job is not to predict the future with certainty. Instead, it gathers evidence, gives each asset a clear score, checks how much risk it carries, and builds a portfolio that avoids putting too much money in one place.")
para(doc, "The main idea is simple: a good decision should not depend on one exciting chart or one headline. The algorithm combines several different clues, then tests whether the approach would still have made sense in earlier periods. A result is treated as more trustworthy only after costs, uncertainty, and possible mistakes have been considered.")

heading(doc, "Research question")
para(doc, "How can a computer program combine company information, price behaviour, market conditions, and risk measurements to create and test a diversified investment portfolio without using information from the future?")

heading(doc, "What the algorithm does")
add_table(doc, ["Step", "Plain language description", "Output"], [
    ("1 Gather", "Collect public historical prices, company accounts, and economic data.", "Clean data set"),
    ("2 Check", "Reject or flag missing, unusual, or unsuitable data.", "Reliable inputs"),
    ("3 Score", "Study each asset from six different viewpoints.", "Scores from 0 to 100"),
    ("4 Combine", "Balance the scores and subtract a penalty for high risk.", "Overall ranking"),
    ("5 Allocate", "Spread money across assets while limiting concentration.", "Target portfolio weights"),
    ("6 Test", "Replay decisions through past periods and include trading costs.", "Evidence and limitations"),
], [0.72, 4.65, 1.05])

heading(doc, "Methodology")
heading(doc, "Data collection and preparation", 2)
para(doc, "The program uses daily historical prices, trading volume, public company financial statements, and broad economic measures such as inflation, unemployment, and interest rate spreads. The standard study period begins in 2015. The S and P 500 is used as a broad market comparison, while short term United States Treasury bills provide a reference for a low risk return.")
para(doc, "Before a calculation is made, the data are checked. An asset generally needs at least three years of history. The program flags too many missing observations and unusually large daily movements. It also handles exchange traded funds differently from companies because funds do not have normal company accounts. If a calculation cannot fairly be applied, the program records a neutral score rather than pretending the data are meaningful.")

heading(doc, "The six evidence modules", 2)
para(doc, "Each module creates a score that is placed on the same 0 to 100 scale. This makes very different kinds of evidence easier to compare. A high score means the evidence is relatively favourable, not that an investment is guaranteed to rise.")
add_table(doc, ["Module", "Question asked", "Student friendly explanation"], [
    ("Business quality", "Is the company financially healthy?", "Looks at profit, cash flow, debt, liquidity, and whether the business is becoming more efficient."),
    ("Valuation", "Does the price seem high or low compared with the business?", "Compares the market price with estimated future cash flows and with similar companies. It uses gradual scoring instead of a harsh pass fail rule."),
    ("Price behaviour", "What has the price been doing?", "Measures momentum, moving averages, and whether recent changes look more like a trend or a short term bounce."),
    ("Market factors", "What kinds of market forces affect this asset?", "Uses a statistical model to see how returns relate to the overall market, company size, value, and momentum factors."),
    ("Risk", "How bad could losses be?", "Studies volatility, the largest past fall, sensitivity to the market, and losses in very bad periods."),
    ("Economic conditions", "What is happening in the wider economy?", "Uses delayed economic information, such as inflation and credit conditions, so the model does not accidentally use news before it was available."),
], [1.2, 1.65, 3.57])

heading(doc, "A closer look at the calculations", 2)
para(doc, "Some calculations have short names, but their purpose can be understood without advanced mathematics. Momentum asks whether an asset has generally been moving upward or downward over one, three, six, and twelve month windows. Volatility measures how much prices jump around. A larger value means a bumpier journey. Drawdown measures the fall from a previous high point to a later low point.")
para(doc, "For a company, the valuation module estimates what future cash the business could generate and converts that idea into a value today. Money expected later is worth less than money available now, so the estimate is reduced by a discount rate. Because every estimate can be wrong, the program also tries several sensible growth and discount rate assumptions rather than trusting one number.")

heading(doc, "Combining the evidence")
para(doc, "The six module scores are normally given equal importance. Their average becomes a starting score. The risk module can then reduce that score using a risk penalty. Finally, assets are ranked from stronger to weaker. This design makes the decision traceable: a reader can see which parts of the evidence helped or hurt an asset.")
para(doc, "A simplified version is: overall score = average of the six evidence scores minus a risk penalty. The real software also keeps warnings, such as missing data or low confidence, beside the score. A score is therefore an organised summary of evidence, not a command to buy.")

heading(doc, "Building the portfolio")
para(doc, "A portfolio is the collection of assets held together. The algorithm chooses weights, which are percentages of the portfolio. It aims for a strong expected reward compared with risk, but it follows safety rules: no short selling, all weights must add to 100 percent, and one asset is normally capped at 40 percent. The program estimates how assets move together, because five assets that all fall at the same time are less diversified than five assets that behave differently.")
para(doc, "It also compares a risk parity approach. Risk parity tries to prevent one risky asset from dominating the portfolio just because it has a high possible return. In practice, diversification does not remove all risk, but it can reduce the damage caused by one poor decision.")

heading(doc, "Testing the method fairly")
para(doc, "The algorithm uses walk forward testing. Imagine standing at a point in the past: the program trains only on information available then, makes a portfolio choice, and checks what happened next. It then moves forward and repeats. The study uses three years of earlier data for learning and one later year for testing, with monthly rebalancing. This is closer to a real decision than mixing past and future data together.")
add_table(doc, ["Fairness check", "Why it matters"], [
    ("No look ahead information", "A historical decision may use only data that would have been published at that time."),
    ("Delayed economic data", "Economic figures are released after the period they describe, so the model uses a one month delay."),
    ("Trading costs", "Every trade is charged 10 basis points, or 0.10 percent, to avoid unrealistic paper profits."),
    ("Benchmark comparison", "Results are compared with the S and P 500 rather than judged in isolation."),
    ("Many measurements", "Return is not enough. The program also reports volatility, Sharpe ratio, maximum drawdown, beta, and other risk measures."),
], [2.2, 4.22])

heading(doc, "How results are judged")
para(doc, "A useful strategy should not only make money in one lucky period. The project compares total and annualised return with volatility and maximum drawdown. It uses the Sharpe ratio to ask how much reward was achieved for each unit of risk. It also uses additional checks to reduce the chance of being fooled by trying many ideas and reporting only the best one.")
para(doc, "The algorithm can be tested under different assumptions, such as different momentum windows, risk limits, and position caps. If a small change in an assumption completely changes the result, the method is fragile and should be treated carefully.")

heading(doc, "Limitations and responsible use")
para(doc, "Historical success does not prove future success. Markets change, data may contain errors, and a model can discover patterns that happened by chance. A backtest may also underestimate the real difficulty of trading, including larger costs, delays, taxes, and the effect of many people trying the same strategy.")
bullet(doc, "The program is for education and research, not personal financial advice.")
bullet(doc, "It does not guarantee profit and can lose money, including during unusual market events.")
bullet(doc, "Students should treat a score as a question to investigate, not a reason to invest immediately.")
bullet(doc, "Any real investment decision should consider personal goals, risk tolerance, rules, and advice from an appropriately qualified adult or professional.")

heading(doc, "Conclusion")
para(doc, "QuantFinance EDU turns a difficult investing question into a repeatable process: collect data, check it, study it from several angles, control risk, and test decisions honestly. The most important lesson is not that a computer can predict markets perfectly. It is that good research is transparent about its evidence, its assumptions, and what it does not know.")

heading(doc, "Glossary")
add_table(doc, ["Term", "Meaning"], [
    ("Algorithm", "A set of instructions a computer follows to solve a problem."),
    ("Backtest", "A simulation that applies a method to past data."),
    ("Diversification", "Spreading investments so one result has less control over the whole portfolio."),
    ("Portfolio weight", "The percentage of the portfolio assigned to one asset."),
    ("Risk", "The possibility and size of unwanted outcomes, especially losses."),
    ("Volatility", "How widely and quickly a price moves up and down."),
    ("Bias", "A mistake in a method that makes results look better or worse than they really are."),
], [1.6, 4.82])

heading(doc, "Empirical Findings")
para(doc, "Historical backtests of this methodology revealed important insights about portfolio construction. For example, testing showed that Hierarchical Risk Parity (HRP) achieved significantly lower turnover (0.009) compared to Sample Covariance (0.056) and Ledoit-Wolf shrinkage (0.056). This empirically supports the claim that HRP produces more stable weight allocations across time.")

para(doc, "However, regarding absolute returns, the advantage of complex models is often challenged by estimation noise. In our primary backtest, HRP produced a higher Sharpe ratio than a simple Equal Weight portfolio, but this difference was rarely statistically significant after multiple-comparison adjustments (Benjamini-Hochberg p = 0.7378). This highlights the sensitivity of financial models to specific data samples and evaluation logic, and aligns with literature suggesting simple 1/N portfolios are difficult to beat consistently.")

heading(doc, "Methodology Note: Debugging & Pipeline Integrity")
para(doc, "During the development of the empirical pipeline, several critical statistical and logic errors were identified and resolved to ensure the integrity of these findings:")
doc.add_paragraph("1. Risk-Free Rate Double Subtraction: Early Sharpe ratios appeared artificially deflated (halving across iterations) because the risk-free rate was being subtracted multiple times in the return calculation pipeline.", style='List Bullet')
doc.add_paragraph("2. Deflated Sharpe Ratio (DSR) Sample Size: The effective number of observations in overlapping Cross-Validated Path blocks was initially overestimated by naively pooling all dates, artificially inflating the DSR. This was corrected to use a single representative test fold size.", style='List Bullet')
doc.add_paragraph("3. Determinism & Random Seeds: Sequential runs with identical parameters yielded wildly divergent results. This was traced to Python's randomized string hashing (used for seeding the synthetic data generator) and unseeded sklearn/scipy estimators. A global deterministic seed architecture was implemented.", style='List Bullet')
doc.add_paragraph("4. Isolated Shrinkage: A fifth method ('LW-Shrinkage-Only Markowitz') was added to the pipeline to causally decompose the performance of Regularised Markowitz into its constituent parts: covariance conditioning versus alpha views.", style='List Bullet')
para(doc, "These rigorous validation steps underscore the difficulty of producing reliable, reproducible quantitative research and the necessity of defensive pipeline engineering.")

heading(doc, "Project sources")
para(doc, "This methodology paper is based on the QuantFinance EDU project configuration, implementation notes, and source code. The project uses public market and economic data. Methods named in the project include company financial analysis, discounted cash flow valuation, price momentum, factor regression, risk measurement, portfolio optimisation, and walk forward backtesting.")

# Footer
for section in doc.sections:
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "QuantFinance EDU  |  Student methodology paper", size=8)

doc.core_properties.title = "How a Quantitative Investing Algorithm Makes Decisions"
doc.core_properties.subject = "Student friendly methodology paper"
doc.core_properties.author = "Swastik"
doc.save(OUT)
print(OUT)
